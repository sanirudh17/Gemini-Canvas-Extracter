"""
Microsoft Word (.docx) Generator for Gemini Canvas Artifacts
Converts Markdown to authentic Microsoft Word documents with:
- Native Word Office Math (OMML) equations
- Native multi-row, multi-column tables with styled headers, borders, and zebra striping
- Precise monospace ASCII diagrams with complete font mapping (Consolas)
- Professional typography with keep-with-next headings and clean list formatting
"""

import io
import os
import re
import functools
from typing import Optional, List, Tuple
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from lxml import etree
import latex2mathml.converter

# ==============================================================================
# OMML (Office Math) Conversion Setup
# ==============================================================================

xslt_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'MML2OMML.XSL')
transform = None
if os.path.exists(xslt_path):
    try:
        xslt_doc = etree.parse(xslt_path)
        transform = etree.XSLT(xslt_doc)
    except Exception as e:
        print(f"Warning: Could not initialize MML2OMML.XSL: {e}")


@functools.lru_cache(maxsize=4096)
def _latex_to_omml_xml(clean: str) -> Optional[bytes]:
    try:
        mathml = latex2mathml.converter.convert(clean)
        tree = etree.fromstring(mathml)
        omml_tree = transform(tree)
        root = omml_tree.getroot()
        return etree.tostring(root) if root is not None else None
    except Exception:
        return None


def latex_to_omml(latex_code: str):
    """
    Converts a LaTeX formula string into a native Word Office Math (OMML) XML element.
    Returns None if conversion fails.
    """
    if not transform:
        return None
    try:
        clean = latex_code.strip()
        # Clean common markdown artifacts inside math
        clean = clean.replace(r'\_', '_').replace(r'\*', '*')
        # Normalize double backslashes
        clean = re.sub(r'\\\\', r'\\', clean)
        xml_bytes = _latex_to_omml_xml(clean)
        if xml_bytes is not None:
            return etree.fromstring(xml_bytes)
        return None
    except Exception:
        return None

def fallback_math_text(math_code: str) -> str:
    """
    High-fidelity Unicode fallback for mathematical formulas when OMML is not applicable.
    """
    s = math_code.strip()
    s = s.replace(r'\_', '_').replace(r'\*', '*')
    
    # Common LaTeX symbol replacements
    replacements = [
        (r'\le', '≤'), (r'\leq', '≤'),
        (r'\ge', '≥'), (r'\geq', '≥'),
        (r'\neq', '≠'), (r'\ne', '≠'),
        (r'\approx', '≈'), (r'\equiv', '≡'),
        (r'\times', '×'), (r'\cdot', '·'),
        (r'\pm', '±'), (r'\mp', '∓'),
        (r'\to', '→'), (r'\rightarrow', '→'), (r'\longrightarrow', '⟶'),
        (r'\leftarrow', '←'), (r'\longleftarrow', '⟵'),
        (r'\implies', '⟹'), (r'\iff', '⟺'),
        (r'\dots', '…'), (r'\cdots', '…'),
        (r'\forall', '∀'), (r'\exists', '∃'),
        (r'\in', '∈'), (r'\notin', '∉'), (r'\subset', '⊂'), (r'\subseteq', '⊆'),
        (r'\cap', '∩'), (r'\cup', '∪'),
        (r'\infty', '∞'),
        (r'\alpha', 'α'), (r'\beta', 'β'), (r'\gamma', 'γ'),
        (r'\delta', 'δ'), (r'\epsilon', 'ε'), (r'\zeta', 'ζ'),
        (r'\eta', 'η'), (r'\theta', 'θ'), (r'\lambda', 'λ'),
        (r'\mu', 'μ'), (r'\nu', 'ν'), (r'\xi', 'ξ'),
        (r'\pi', 'π'), (r'\rho', 'ρ'), (r'\sigma', 'σ'),
        (r'\tau', 'τ'), (r'\phi', 'φ'), (r'\chi', 'χ'),
        (r'\psi', 'ψ'), (r'\omega', 'ω'),
        (r'\Delta', 'Δ'), (r'\Sigma', 'Σ'), (r'\Omega', 'Ω'),
        (r'\Theta', 'Θ'), (r'\Lambda', 'Λ'), (r'\Phi', 'Φ'),
    ]
    for pattern, rep in replacements:
        s = s.replace(pattern, rep)

    # Fractions: \frac{a}{b} -> (a / b)
    s = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', s)
    # Square root: \sqrt{x} -> √(x)
    s = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', s)

    # Convert common superscripts
    sup_map = {'0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴', '5': '⁵',
               '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹', '+': '⁺', '-': '⁻',
               '=': '⁼', '(': '⁽', ')': '⁾', 'n': 'ⁿ', 'm': 'ᵐ', 'i': 'ⁱ'}
    def sup_repl(m):
        content = m.group(1)
        res = ''.join(sup_map.get(c, c) for c in content)
        return res
    s = re.sub(r'\^\{([^}]+)\}', sup_repl, s)
    s = re.sub(r'\^([0-9m-n\+\-])', lambda m: sup_map.get(m.group(1), m.group(1)), s)

    # Convert common subscripts
    sub_map = {'0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄', '5': '₅',
               '6': '₆', '7': '₇', '8': '₈', '9': '₉', '+': '₊', '-': '₋',
               '=': '₌', '(': '₍', ')': '₎', 'a': 'ₐ', 'e': 'ₑ', 'h': 'ₕ',
               'i': 'ᵢ', 'j': 'ⱼ', 'k': 'ₖ', 'l': 'ₗ', 'm': 'ₘ', 'n': 'ₙ',
               'o': 'ₒ', 'p': 'ₚ', 'r': 'ᵣ', 's': 'ₛ', 't': 'ₜ', 'u': 'ᵤ',
               'v': 'ᵥ', 'x': 'ₓ'}
    def sub_repl(m):
        content = m.group(1)
        res = ''.join(sub_map.get(c, c) for c in content)
        return res
    s = re.sub(r'_\{([^}]+)\}', sub_repl, s)
    s = re.sub(r'_([0-9a-z])', lambda m: sub_map.get(m.group(1), m.group(1)), s)

    # Clean leftover backslashes before words
    s = re.sub(r'\\([a-zA-Z]+)', r'\1', s)
    return s


# ==============================================================================
# Typography & XML Helper Functions
# ==============================================================================

def set_run_font(
    run,
    font_name: str,
    font_size: Optional[Pt] = None,
    bold: bool = False,
    italic: bool = False,
    color_rgb: Optional[RGBColor] = None
):
    """Sets run font properties and ensures complete w:rFonts mapping for Windows Word."""
    run.font.name = font_name
    if font_size:
        run.font.size = font_size
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}" w:eastAsia="{font_name}"/>'
    )
    rPr.append(rFonts)


def add_inline_formatted_text(
    p,
    text: str,
    base_font: str = 'Inter',
    base_size: Pt = Pt(10),
    base_color: RGBColor = RGBColor(15, 23, 42)
):
    """
    Parses inline markdown (**bold**, *italic*, `code`, and $math$) and adds
    corresponding native runs and OMML math elements to the paragraph.
    """
    # Tokenize text by markdown delimiters
    tokens = re.split(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\$[^\$]+?\$)', text)

    for token in tokens:
        if not token:
            continue

        if token.startswith('**') and token.endswith('**') and len(token) > 4:
            subtext = token[2:-2]
            if '$' in subtext:
                add_inline_formatted_text(p, subtext, base_font=base_font, base_size=base_size, base_color=base_color)
            else:
                r = p.add_run(subtext)
                set_run_font(r, base_font, base_size, bold=True, color_rgb=base_color)

        elif token.startswith('*') and token.endswith('*') and len(token) > 2:
            r = p.add_run(token[1:-1])
            set_run_font(r, base_font, base_size, italic=True, color_rgb=base_color)

        elif token.startswith('`') and token.endswith('`') and len(token) > 2:
            r = p.add_run(token[1:-1])
            set_run_font(r, 'Consolas', Pt(base_size.pt * 0.92), color_rgb=RGBColor(37, 99, 235))

        elif token.startswith('$') and token.endswith('$') and len(token) > 2:
            math_code = token[1:-1]
            omml = latex_to_omml(math_code)
            if omml is not None:
                p._element.append(omml)
            else:
                clean = fallback_math_text(math_code)
                r = p.add_run(clean)
                set_run_font(r, 'Cambria Math', base_size, italic=True, color_rgb=RGBColor(30, 41, 59))
        else:
            r = p.add_run(token)
            set_run_font(r, base_font, base_size, color_rgb=base_color)


# ==============================================================================
# Markdown Block Parser
# ==============================================================================

def parse_markdown_blocks(text: str) -> List[Tuple]:
    """
    Parses a markdown string into structured semantic blocks:
    - ('heading', level, text)
    - ('table', list_of_rows)
    - ('code', lang, code_text)
    - ('display_math', math_text)
    - ('blockquote', text)
    - ('hr',)
    - ('list_item', 'bullet'|'number', text)
    - ('paragraph', text)
    """
    lines = text.replace('\r\n', '\n').split('\n')
    blocks = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Empty lines
        if not stripped:
            i += 1
            continue

        # 1. Fenced Code Block (```)
        if stripped.startswith('```'):
            lang = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            if i < n:  # Skip closing ```
                i += 1
            blocks.append(('code', lang, '\n'.join(code_lines)))
            continue

        # 2. Display Math ($$)
        if stripped.startswith('$$'):
            if len(stripped) > 2 and stripped.endswith('$$'):
                blocks.append(('display_math', stripped[2:-2].strip()))
                i += 1
                continue
            math_lines = []
            i += 1
            while i < n and not lines[i].strip().endswith('$$'):
                math_lines.append(lines[i])
                i += 1
            if i < n:
                math_lines.append(lines[i].strip()[:-2])
                i += 1
            blocks.append(('display_math', '\n'.join(math_lines).strip()))
            continue

        # 3. Headings (#)
        h_match = re.match(r'^(#{1,6})\s+(.+)$', stripped)
        if h_match:
            level = len(h_match.group(1))
            h_text = h_match.group(2).strip()
            blocks.append(('heading', level, h_text))
            i += 1
            continue

        # 4. Horizontal Rule (---, ***, ___)
        if re.match(r'^(-{3,}|\*{3,}|_{3,})$', stripped):
            blocks.append(('hr',))
            i += 1
            continue

        # 5. GFM Table (| Header | ... |)
        if '|' in stripped and i + 1 < n and re.match(r'^\s*\|?\s*[-:|\s]+\s*\|?\s*$', lines[i+1]):
            table_lines = [stripped]
            i += 2  # Skip header and separator
            while i < n and '|' in lines[i].strip() and lines[i].strip():
                table_lines.append(lines[i].strip())
                i += 1

            parsed_rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip('|').split('|')]
                parsed_rows.append(cells)
            blocks.append(('table', parsed_rows))
            continue

        # 6. Blockquote (>)
        if stripped.startswith('>'):
            bq_lines = []
            while i < n and (lines[i].strip().startswith('>') or (lines[i].strip() and bq_lines)):
                if lines[i].strip().startswith('>'):
                    bq_lines.append(lines[i].strip().lstrip('> ').strip())
                    i += 1
                else:
                    break
            blocks.append(('blockquote', ' '.join(bq_lines)))
            continue

        # 7. List Items
        bullet_match = re.match(r'^([-*+])\s+(.+)$', stripped)
        if bullet_match:
            blocks.append(('list_item', 'bullet', bullet_match.group(2).strip()))
            i += 1
            continue

        num_match = re.match(r'^(\d+)\.\s+(.+)$', stripped)
        if num_match:
            blocks.append(('list_item', 'number', num_match.group(2).strip()))
            i += 1
            continue

        # 8. Paragraph (accumulate lines until the next block delimiter)
        p_lines = [stripped]
        i += 1
        while i < n:
            next_s = lines[i].strip()
            if not next_s:
                break
            if next_s.startswith('```') or next_s.startswith('$$') or next_s.startswith('#') or \
               re.match(r'^(-{3,}|\*{3,}|_{3,})$', next_s) or \
               ('|' in next_s and i + 1 < n and re.match(r'^\s*\|?\s*[-:|\s]+\s*\|?\s*$', lines[i+1])) or \
               next_s.startswith('>') or re.match(r'^([-*+]|\d+\.)\s+', next_s):
                break
            p_lines.append(next_s)
            i += 1
        blocks.append(('paragraph', ' '.join(p_lines)))

    return blocks


# ==============================================================================
# Document Block Builders
# ==============================================================================

def add_header_card(doc, title: str, metadata: Optional[dict] = None):
    """Adds a modern executive header card to the top of the Word document."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.9)

    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    tcPr.append(shd)

    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="160" w:type="dxa"/>
            <w:bottom w:w="160" w:type="dxa"/>
            <w:left w:w="200" w:type="dxa"/>
            <w:right w:w="200" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run("GEMINI CANVAS ARTIFACT")
    set_run_font(r1, 'Plus Jakarta Sans', Pt(8.5), bold=True, color_rgb=RGBColor(37, 99, 235))

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(4)
    r2 = p2.add_run(title)
    set_run_font(r2, 'Plus Jakarta Sans', Pt(15), bold=True, color_rgb=RGBColor(15, 23, 42))

    if metadata:
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(0)
        meta_parts = []
        if metadata.get('share_id'):
            meta_parts.append(f"Share ID: {metadata['share_id']}")
        if metadata.get('word_count'):
            meta_parts.append(f"{metadata['word_count']:,} words")
        if metadata.get('ascii_count'):
            meta_parts.append(f"{metadata['ascii_count']} diagrams")
        r3 = p3.add_run("  •  ".join(meta_parts))
        set_run_font(r3, 'Inter', Pt(8.5), color_rgb=RGBColor(100, 116, 139))

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(12)


def add_ascii_diagram(doc, code_text: str, light_theme: bool = True):
    """
    Inserts a monospace ASCII diagram block with complete font definition
    and keep-together properties.
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.9)

    # CantSplit on table row
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    bg_hex = "F8FAFC" if light_theme else "090D16"
    border_hex = "CBD5E1" if light_theme else "1E293B"
    text_color = RGBColor(15, 23, 42) if light_theme else RGBColor(56, 189, 248)

    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    tcPr.append(shd)

    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="3B82F6"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="140" w:type="dxa"/>
            <w:right w:w="140" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    lines = code_text.split('\n')
    for idx, line in enumerate(lines):
        r = p.add_run(line)
        set_run_font(r, 'Consolas', Pt(8.0), color_rgb=text_color)
        if idx < len(lines) - 1:
            p.add_run('\n')

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)


def add_native_word_table(doc, rows_data: List[List[str]]):
    """
    Constructs a native Microsoft Word table from parsed row data.
    Features styled headers, thin borders, zebra striping, and full cell inline formatting.
    """
    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)

    tbl = doc.add_table(rows=num_rows, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Set table borders
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    for r_idx, row_data in enumerate(rows_data):
        row = tbl.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        is_header = (r_idx == 0)
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            val = row_data[c_idx] if c_idx < len(row_data) else ""

            tcPr = cell._tc.get_or_add_tcPr()
            bg_color = "F1F5F9" if is_header else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
            tcPr.append(shd)

            tcMar = parse_xml(f'''
                <w:tcMar {nsdecls("w")}>
                    <w:top w:w="120" w:type="dxa"/>
                    <w:bottom w:w="120" w:type="dxa"/>
                    <w:left w:w="140" w:type="dxa"/>
                    <w:right w:w="140" w:type="dxa"/>
                </w:tcMar>
            ''')
            tcPr.append(tcMar)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15

            add_inline_formatted_text(
                p, val,
                base_font='Inter',
                base_size=Pt(9.0 if is_header else 8.5),
                base_color=RGBColor(15, 23, 42)
            )

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)


# ==============================================================================
# Main Generator Entrypoint
# ==============================================================================

def generate_docx(
    title: str,
    markdown_content: str,
    ascii_light_theme: bool = True,
    metadata: Optional[dict] = None
) -> bytes:
    """
    Main entrypoint: Generates a complete, authentic Microsoft Word document (.docx)
    from markdown content with native tables, OMML equations, and monospace diagrams.
    """
    doc = docx.Document()

    # Standard 0.8 inch margins for balanced page distribution
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

        # Header and Footer
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Gemini Canvas Export")
        set_run_font(frun, 'Inter', Pt(8.0), color_rgb=RGBColor(148, 163, 184))

    # Add Executive Header Block
    add_header_card(doc, title, metadata)

    # Parse markdown into structured semantic blocks
    blocks = parse_markdown_blocks(markdown_content)

    for b in blocks:
        btype = b[0]

        if btype == 'heading':
            level = min(b[1], 4)
            text = b[2]
            h = doc.add_heading(level=level)
            h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            add_inline_formatted_text(
                h, text,
                base_font='Plus Jakarta Sans',
                base_size=Pt(16 if level == 1 else (13.5 if level == 2 else 11.5)),
                base_color=RGBColor(15, 23, 42)
            )

        elif btype == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/></w:pBdr>'
            )
            p._p.get_or_add_pPr().append(pBdr)

        elif btype == 'table':
            add_native_word_table(doc, b[1])

        elif btype == 'code':
            add_ascii_diagram(doc, b[2], light_theme=ascii_light_theme)

        elif btype == 'display_math':
            math_code = b[1]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            omml = latex_to_omml(math_code)
            if omml is not None:
                p._element.append(omml)
            else:
                clean = fallback_math_text(math_code)
                r = p.add_run(clean)
                set_run_font(r, 'Cambria Math', Pt(11), italic=True, color_rgb=RGBColor(30, 41, 59))

        elif btype == 'blockquote':
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.9)

            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F9FF"/>')
            tcPr.append(shd)

            borders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>
                    <w:top w:val="none"/>
                    <w:right w:val="none"/>
                    <w:bottom w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(borders)

            tcMar = parse_xml(f'''
                <w:tcMar {nsdecls("w")}>
                    <w:top w:w="100" w:type="dxa"/>
                    <w:bottom w:w="100" w:type="dxa"/>
                    <w:left w:w="140" w:type="dxa"/>
                    <w:right w:w="140" w:type="dxa"/>
                </w:tcMar>
            ''')
            tcPr.append(tcMar)

            bp = cell.paragraphs[0]
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.space_after = Pt(0)
            bp.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(
                bp, b[1],
                base_font='Inter',
                base_size=Pt(9.5),
                base_color=RGBColor(3, 105, 161)
            )

            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(6)

        elif btype == 'list_item':
            ltype = b[1]
            text = b[2]
            style_name = 'List Bullet' if ltype == 'bullet' else 'List Number'
            lp = doc.add_paragraph(style=style_name)
            lp.paragraph_format.space_before = Pt(0)
            lp.paragraph_format.space_after = Pt(3)
            lp.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(lp, text, base_font='Inter', base_size=Pt(9.5))

        elif btype == 'paragraph':
            text = b[1]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatted_text(p, text, base_font='Inter', base_size=Pt(9.5))

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
